from pathlib import Path
from html.parser import HTMLParser
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
BASE_URL = "https://aprendanetstorage.github.io/questoesaprendanet"

NOME_BONITO = {
    "cinematica": "Cinematica",
    "leis_newton": "Leis de Newton",
    "energia": "Energia",
    "quantidade_movimento": "Quantidade de Movimento",
    "torque": "Torque e Equilibrio",
    "termologia": "Termologia",
    "ondulatoria": "Ondulatoria",
    "optica": "Optica",
    "eletricidade": "Eletricidade",
    "fisica": "Fisica (Geral)",
}

SUBTOPICOS = {
    "cinematica": "MRU; MRUV; queda livre; lancamento obliquo; lancamento horizontal; graficos espaco-tempo; graficos velocidade-tempo",
    "leis_newton": "1a lei de Newton; 2a lei de Newton; 3a lei de Newton; plano inclinado; forca de atrito; forca normal; dinamica",
    "energia": "energia cinetica; energia potencial gravitacional; trabalho de uma forca; potencia mecanica; conservacao de energia; rendimento de maquinas",
    "quantidade_movimento": "impulso; teorema impulso-variacao; colisao elastica; colisao inelastica; conservacao do momento linear",
    "torque": "momento de forca; equilibrio de corpos extensos; alavancas; centro de massa; condicoes de equilibrio estatico",
    "termologia": "temperatura; escalas termometricas; dilatacao termica; calorimetria; mudanca de fase; lei dos gases; ciclo de Carnot",
    "ondulatoria": "classificacao de ondas; comprimento de onda; frequencia; periodo; efeito Doppler; ressonancia; ondas estacionarias",
    "optica": "reflexao; refracao; lentes convergentes; lentes divergentes; espelhos planos; espelhos esfericos; indice de refracao; cor",
    "eletricidade": "carga eletrica; campo eletrico; corrente eletrica; resistencia; lei de Ohm; circuito serie e paralelo; magnetismo; inducao eletromagnetica",
    "fisica": "grandezas fisicas; medicao; fenomenos fisicos; questoes ENEM diversas",
}

KEYWORDS = {
    "cinematica": "velocidade, aceleracao, deslocamento, trajetoria, MRU, MRUV, queda livre, lancamento, posicao, movimento",
    "leis_newton": "forca, Newton, atrito, normal, peso, inercia, 2a lei, 3a lei, acao e reacao, dinamica, resultante",
    "energia": "energia cinetica, energia potencial, trabalho, potencia, conservacao de energia, joule, watt, rendimento, eficiencia",
    "quantidade_movimento": "impulso, quantidade de movimento, colisao, momento linear, choque elastico, choque inelastico, explosao",
    "torque": "torque, momento de forca, equilibrio, alavanca, centro de massa, estatica, braço da forca",
    "termologia": "temperatura, calor, dilatacao, gas, pressao, volume, termodinamica, lei dos gases, entropia, calorimetria",
    "ondulatoria": "onda, frequencia, comprimento de onda, amplitude, periodo, som, acustica, doppler, ressonancia, onda estacionaria",
    "optica": "luz, lente, espelho, refracao, reflexao, difracao, prisma, optica, convergente, divergente, cor, indice de refracao",
    "eletricidade": "corrente eletrica, tensao, resistencia, circuito, ohm, capacitor, campo eletrico, eletromagnetismo, magnetismo, induçao",
    "fisica": "fisica, ENEM, fenomeno fisico, grandeza, medida, questao diversa",
}

TITULO_LIXO = re.compile(
    r"(sistemas num[ée]ricos|lista\s*0[0-9]|problema\s*\d+|document|lista\s*0[1-9])",
    re.IGNORECASE,
)


class TitleBodyParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.title = ""
        self.body_text = []
        self._in_title = False
        self._in_style = False
        self._in_script = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        elif tag == "style":
            self._in_style = True
        elif tag == "script":
            self._in_script = True

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        elif tag == "style":
            self._in_style = False
        elif tag == "script":
            self._in_script = False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._in_style and not self._in_script:
            stripped = data.strip()
            if stripped:
                self.body_text.append(stripped)

    def get_enunciado(self):
        texto = " ".join(self.body_text)
        # pega primeiras 180 chars de texto real
        texto = re.sub(r"\s+", " ", texto).strip()
        return texto[:180] + "..." if len(texto) > 180 else texto


def titulo_limpo(titulo_raw: str, materia: str, num: int) -> str:
    titulo = titulo_raw.strip()
    # remove encoding quebrado
    titulo = titulo.replace("�", "").strip()
    # se for lixo, gera título padrao
    if not titulo or TITULO_LIXO.search(titulo):
        return f"Questao {num:03d} — {NOME_BONITO[materia]}"
    # remove sufixo " — Cinemática" duplicado se existir
    titulo = re.sub(r"\s*—\s*.+$", "", titulo).strip()
    return titulo or f"Questao {num:03d} — {NOME_BONITO[materia]}"


def descricao(enunciado: str, materia: str) -> str:
    if enunciado and len(enunciado) > 30:
        return enunciado
    return f"Questao do ENEM sobre {NOME_BONITO[materia].lower()}."


# ── Coleta dados ──────────────────────────────────────────────────────────────

rows = []
materias = sorted(
    [d for d in BASE.iterdir() if d.is_dir() and d.name in NOME_BONITO],
    key=lambda x: x.name,
)

for pasta in materias:
    questoes = sorted(pasta.glob("questao_*.html"))
    for i, q in enumerate(questoes, start=1):
        txt = q.read_text(encoding="utf-8", errors="replace")
        p = TitleBodyParser()
        p.feed(txt)
        titulo = titulo_limpo(p.title, pasta.name, i)
        enunc = descricao(p.get_enunciado(), pasta.name)
        url = f"{BASE_URL}/{pasta.name}/{q.name}"
        rows.append({
            "Materia": "Fisica",
            "Topico Principal": NOME_BONITO[pasta.name],
            "Subtopicos Relacionados": SUBTOPICOS[pasta.name],
            "Palavras-chave": KEYWORDS[pasta.name],
            "Descricao Natural": enunc,
            "Titulo da Questao": titulo,
            "Link da Questao": url,
        })

print(f"Linhas coletadas: {len(rows)}")

# ── Gera Excel ────────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Questoes"

HEADERS = [
    "Materia",
    "Topico Principal",
    "Subtopicos Relacionados",
    "Palavras-chave",
    "Descricao Natural",
    "Titulo da Questao",
    "Link da Questao",
]

WIDTHS = [12, 22, 55, 55, 65, 35, 70]

header_fill = PatternFill("solid", start_color="2C5F8A", end_color="2C5F8A")
header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

cell_font = Font(name="Arial", size=10)
cell_align = Alignment(vertical="top", wrap_text=True)

alt_fill = PatternFill("solid", start_color="EDF3FA", end_color="EDF3FA")

for col, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = border
    ws.column_dimensions[get_column_letter(col)].width = width

ws.row_dimensions[1].height = 30

for r, row in enumerate(rows, start=2):
    is_alt = (r % 2 == 0)
    for c, key in enumerate(HEADERS, start=1):
        # mapeia nome da coluna para chave do dict
        key_map = {
            "Materia": "Materia",
            "Topico Principal": "Topico Principal",
            "Subtopicos Relacionados": "Subtopicos Relacionados",
            "Palavras-chave": "Palavras-chave",
            "Descricao Natural": "Descricao Natural",
            "Titulo da Questao": "Titulo da Questao",
            "Link da Questao": "Link da Questao",
        }
        value = row[key_map[key]]
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = cell_font
        cell.alignment = cell_align
        cell.border = border
        if is_alt:
            cell.fill = alt_fill

    ws.row_dimensions[r].height = 45

ws.freeze_panes = "A2"

out = BASE.parent / "Questoes_ENEM_Fisica.xlsx"
wb.save(out)
print(f"Salvo em: {out}")
